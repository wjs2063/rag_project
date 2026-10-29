"""Bounded-memory preparation of zipped SHP/DBF data into a disk-backed staging DB."""

import hashlib
import heapq
import json
import shutil
import sqlite3
import tempfile
from collections import Counter
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

import shapefile
from openpyxl import load_workbook
from pyproj import Transformer

from poi_service.catalog import Catalog, catalog_path
from poi_service.geo import point_wkt


def _field(row: dict, name: str | None) -> str:
    if not name:
        return ""
    value = row.get(name.upper())
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@contextmanager
def _members(archive: Path, names: dict[str, str]) -> Iterator[dict[str, Path]]:
    """Extract only declared ZIP members to a short-lived directory."""
    with tempfile.TemporaryDirectory(prefix="poi-region-") as tmp, ZipFile(archive) as source:
        result: dict[str, Path] = {}
        for role, member in names.items():
            if member not in source.namelist():
                raise ValueError(f"missing ZIP member: {archive}:{member}")
            target = Path(tmp) / f"{role}{Path(member).suffix.lower()}"
            with source.open(member) as input_file, target.open("wb") as output_file:
                shutil.copyfileobj(input_file, output_file, length=1024 * 1024)
            result[role] = target
        yield result


def _dbf_rows(archive: Path, member: str, encoding: str) -> Iterator[dict]:
    with _members(archive, {"table": member}) as paths:
        reader = shapefile.Reader(dbf=str(paths["table"]), encoding=encoding, encodingErrors="replace")
        try:
            for record in reader.iterRecords():
                yield record.as_dict()
        finally:
            reader.close()


def _xlsx_rows(archive: Path, member: str, sheet_name: str | None, header_row: int) -> Iterator[dict]:
    with _members(archive, {"table": member}) as paths:
        book = load_workbook(paths["table"], read_only=True, data_only=True)
        try:
            sheet = book[sheet_name] if sheet_name else book.active
            rows = sheet.iter_rows(values_only=True)
            for _ in range(header_row - 1):
                next(rows)
            headers = [str(cell).strip().upper() if cell is not None else "" for cell in next(rows)]
            if not any(headers):
                raise ValueError(f"empty XLSX header: {archive}:{member}")
            for values in rows:
                yield dict(zip(headers, values))
        finally:
            book.close()


def _resolve(base: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else base / path


def _schema(connection: sqlite3.Connection) -> None:
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS pois (nf_id TEXT PRIMARY KEY, payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS aux (source TEXT NOT NULL, join_id TEXT NOT NULL, payload TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS aux_lookup ON aux(source, join_id);
        CREATE TABLE IF NOT EXISTS embeddings (nf_id TEXT PRIMARY KEY, vector TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS reports (region TEXT PRIMARY KEY, payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
    """)


def _load_aux(connection: sqlite3.Connection, base: Path, sources: list[dict], region: str) -> Counter:
    counts: Counter = Counter()
    for spec in sources:
        archive = _resolve(base, spec["zip"])
        member = spec["member"]
        if member.lower().endswith(".xlsx"):
            rows = _xlsx_rows(archive, member, spec.get("sheet"), spec.get("header_row", 1))
        else:
            rows = _dbf_rows(archive, member, spec.get("encoding", "cp949"))
        for raw_record in rows:
            record = {str(key).upper(): value for key, value in raw_record.items()}
            counts[f"aux_{spec['name']}_rows"] += 1
            key = _field(record, spec["key_field"])
            if not key:
                counts[f"aux_{spec['name']}_missing_key"] += 1
                continue
            mapped = {}
            for dest, source in spec["fields"].items():
                if isinstance(source, list):
                    values = [_field(record, name) for name in source]
                    mapped[dest] = " ".join(filter(None, values)) if dest == "alias" else next(
                        (value for value in values if value), ""
                    )
                else:
                    mapped[dest] = _field(record, source)
            if not any(mapped.values()):
                counts[f"aux_{spec['name']}_empty"] += 1
                continue
            if any("\ufffd" in value for value in mapped.values()):
                counts[f"aux_{spec['name']}_replacement_text"] += 1
            connection.execute(
                "INSERT INTO aux VALUES (?,?,?)", (spec["name"], key, json.dumps(mapped, ensure_ascii=False))
            )
        connection.commit()
    return counts


def _lookup(connection: sqlite3.Connection, source: str, key: str) -> list[dict]:
    if not key:
        return []
    return [
        json.loads(row[0])
        for row in connection.execute("SELECT payload FROM aux WHERE source=? AND join_id=?", (source, key))
    ]


def _limit_utf8(value: str, max_bytes: int) -> str:
    return value.encode("utf-8")[:max_bytes].decode("utf-8", errors="ignore")


def _prepare_region(
    connection: sqlite3.Connection, base: Path, spec: dict, catalog: Catalog, source_version: str
) -> Counter:
    counts: Counter = Counter()
    region = spec["name"]
    if spec.get("exclude", False):
        archive = _resolve(base, spec["zip"])
        with ZipFile(archive) as source, source.open(spec["dbf_member"]) as table:
            header = table.read(32)
        counts["excluded_region_rows"] = int.from_bytes(header[4:8], "little")
        return counts
    connection.execute("DELETE FROM aux WHERE source NOT LIKE 'global:%'")
    connection.commit()
    auxiliaries = spec.get("global_auxiliary", []) + spec.get("auxiliary", [])
    counts.update(_load_aux(connection, base, spec.get("auxiliary", []), region))
    archive = _resolve(base, spec["zip"])
    mapping = spec["fields"]
    members = {"shp": spec["shp_member"], "dbf": spec["dbf_member"]}
    if spec.get("shx_member"):
        members["shx"] = spec["shx_member"]
    transformer = Transformer.from_crs(spec.get("source_crs", "EPSG:5179"), "EPSG:4326", always_xy=True)
    with _members(archive, members) as paths:
        reader = shapefile.Reader(
            shp=str(paths["shp"]),
            dbf=str(paths["dbf"]),
            shx=str(paths["shx"]) if "shx" in paths else None,
            encoding=spec.get("encoding", "cp949"),
            encodingErrors="replace",
        )
        try:
            for record, shape in zip(reader.iterRecords(), reader.iterShapes(), strict=True):
                counts["source_rows"] += 1
                row = {str(key).upper(): value for key, value in record.as_dict().items()}
                if any("\ufffd" in value for value in row.values() if isinstance(value, str)):
                    counts["source_replacement_text"] += 1
                nf_id = _field(row, mapping["nf_id"])
                poi_id = _field(row, mapping.get("poiid"))
                if not nf_id:
                    counts["missing_nf_id"] += 1
                    continue
                values = {dest: _field(row, source) for dest, source in mapping.items()}
                for aux in auxiliaries:
                    key = poi_id if aux.get("join_on", "NF_ID").upper() == "POIID" else nf_id
                    matches = _lookup(connection, aux["name"], key)
                    if not matches:
                        counts[f"missing_join_{aux['name']}"] += 1
                    for match in matches:
                        for dest, value in match.items():
                            if value:
                                if dest == "alias":
                                    values.setdefault("aliases", []).append(value)
                                elif not values.get(dest):
                                    values[dest] = value
                category_id = values.get("category_id", "")
                category = catalog.categories.get(category_id)
                if not category:
                    counts["unknown_category"] += 1
                    counts[f"unknown_code:{category_id or '<empty>'}"] += 1
                    continue
                if not category.allowed:
                    counts["excluded_category"] += 1
                    continue
                if not values.get("name"):
                    counts["missing_name"] += 1
                    continue
                if not shape.points:
                    counts["missing_coordinate"] += 1
                    continue
                try:
                    lon, lat = transformer.transform(*shape.points[0][:2])
                    geo = point_wkt(lon, lat)
                except (ValueError, OverflowError):
                    counts["invalid_coordinate"] += 1
                    continue
                # The initial product covers South Korea. A separate flag can exclude North Korea by region.
                if not (124 <= lon <= 132 and 33 <= lat <= 39.9):
                    counts["outside_service_bounds"] += 1
                    continue
                aliases = " ".join(dict.fromkeys(values.get("aliases", [])))
                name = _limit_utf8(values["name"], 512)
                address = _limit_utf8(values.get("address", ""), 1024)
                search_text = _limit_utf8(f"{name} {aliases}", 3072)
                payload = {
                    "nf_id": _limit_utf8(nf_id, 64),
                    "name": name,
                    "aliases": _limit_utf8(aliases, 2048),
                    "search_text": search_text,
                    "address": address,
                    "category_id": category_id,
                    "category_name": _limit_utf8(category.name, 256),
                    "category_parent_id": category.parent_id or "",
                    "longitude": lon,
                    "latitude": lat,
                    "geo": geo,
                    "source_region": region,
                    "source_version": source_version,
                    "updated_at": datetime.now(UTC).isoformat(timespec="seconds"),
                }
                try:
                    connection.execute(
                        "INSERT INTO pois VALUES (?,?)",
                        (payload["nf_id"], json.dumps(payload, ensure_ascii=False)),
                    )
                    counts["selected_rows"] += 1
                    counts[f"selected_category:{category_id}"] += 1
                except sqlite3.IntegrityError:
                    counts["duplicate_nf_id"] += 1
                if counts["source_rows"] % 10_000 == 0:
                    connection.commit()
        finally:
            reader.close()
    connection.commit()
    if checks := spec.get("xlsx_check"):
        if isinstance(checks, dict):
            checks = [checks]
        for check in checks:
            with _members(_resolve(base, check["zip"]), {"xlsx": check["member"]}) as paths:
                book = load_workbook(paths["xlsx"], read_only=True, data_only=True)
                try:
                    sheet = book[check["sheet"]] if check.get("sheet") else book.active
                    counts["xlsx_rows"] += sum(1 for _ in sheet.rows) - int(check.get("header_rows", 1))
                finally:
                    book.close()
        if counts["xlsx_rows"] != counts["source_rows"]:
            counts["xlsx_count_mismatch"] = counts["xlsx_rows"] - counts["source_rows"]
    return counts


def prepare(manifest_path: Path, output_db: Path) -> dict:
    """Normalize one region at a time; retain compact SQLite staging and discrepancy reports."""
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    catalog_file, version = catalog_path(manifest_path)
    catalog = Catalog.from_csv(catalog_file, version)
    if not manifest.get("regions") or not manifest.get("source_version"):
        raise ValueError("source_version and regions are required")
    if output_db.exists():
        raise FileExistsError(output_db)
    output_db.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "catalog_version": version,
        "catalog_category_count": len(catalog.categories),
        "catalog_allowed_count": sum(category.allowed for category in catalog.categories.values()),
        "regions": {},
    }
    with sqlite3.connect(output_db) as connection:
        connection.execute("PRAGMA cache_size=-65536")
        connection.execute("PRAGMA temp_store=FILE")
        _schema(connection)
        global_auxiliary = manifest.get("global_auxiliary", [])
        if global_auxiliary:
            if any(not source["name"].startswith("global:") for source in global_auxiliary):
                raise ValueError("global auxiliary names must start with 'global:'")
            report["global_auxiliary"] = dict(
                _load_aux(connection, manifest_path.parent, global_auxiliary, "global")
            )
        for spec in manifest["regions"]:
            spec["global_auxiliary"] = global_auxiliary
            counts = _prepare_region(
                connection, manifest_path.parent, spec, catalog, manifest["source_version"]
            )
            report["regions"][spec["name"]] = dict(counts)
            connection.execute("INSERT INTO reports VALUES (?,?)", (spec["name"], json.dumps(dict(counts))))
            connection.commit()
        report["selected_total"] = connection.execute("SELECT COUNT(*) FROM pois").fetchone()[0]
    output_db.with_suffix(".report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def sample_staging(source_db: Path, output_db: Path, regions: list[str], per_region: int = 2000,
                   seed: str = "poi-sample-v1") -> dict:
    """Select a repeatable bounded sample using the smallest seeded hashes per region."""
    if not regions or len(set(regions)) != len(regions) or per_region < 1:
        raise ValueError("provide distinct regions and a positive per-region sample size")
    if output_db.exists():
        raise FileExistsError(output_db)
    output_db.parent.mkdir(parents=True, exist_ok=True)
    heaps: dict[str, list[tuple[int, str, str]]] = {region: [] for region in regions}
    with sqlite3.connect(source_db) as source:
        for nf_id, payload in source.execute("SELECT nf_id,payload FROM pois"):
            row = json.loads(payload)
            region = row["source_region"]
            if region not in heaps:
                continue
            priority = int.from_bytes(
                hashlib.blake2b(f"{seed}:{nf_id}".encode(), digest_size=8).digest(), "big"
            )
            entry = (-priority, nf_id, payload)
            heap = heaps[region]
            if len(heap) < per_region:
                heapq.heappush(heap, entry)
            elif entry > heap[0]:
                heapq.heapreplace(heap, entry)
        counts = {region: len(heap) for region, heap in heaps.items()}
        if any(count != per_region for count in counts.values()):
            raise ValueError(f"insufficient eligible POIs for requested sample: {counts}")
        with sqlite3.connect(output_db) as target:
            _schema(target)
            for heap in heaps.values():
                target.executemany(
                    "INSERT INTO pois(nf_id,payload) VALUES (?,?)",
                    [(nf_id, payload) for _, nf_id, payload in heap],
                )
            for row in source.execute("SELECT key,value FROM metadata"):
                target.execute("INSERT INTO metadata VALUES (?,?)", row)
    selected_ids = sorted(nf_id for heap in heaps.values() for _, nf_id, _ in heap)
    report = {
        "seed": seed,
        "per_region": per_region,
        "regions": counts,
        "selected_total": sum(counts.values()),
        "selected_ids_sha256": hashlib.sha256("\n".join(selected_ids).encode()).hexdigest(),
    }
    output_db.with_suffix(".report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def inspect_zip(archive: Path, headers: bool = False) -> list[dict]:
    """List members, DBF counts and optionally DBF/XLSX column headers."""
    members = []
    with ZipFile(archive) as source:
        for info in source.infolist():
            if info.is_dir():
                continue
            entry = {"name": info.filename, "bytes": info.file_size}
            if info.filename.lower().endswith(".dbf"):
                with source.open(info) as stream:
                    header = stream.read(32)
                    if headers and len(header) == 32:
                        header_length = int.from_bytes(header[8:10], "little")
                        descriptors = stream.read(max(0, header_length - 32))
                        entry["fields"] = [
                            descriptors[offset:offset + 11].split(b"\0", 1)[0].decode("cp949", "replace").strip()
                            for offset in range(0, len(descriptors) - 1, 32)
                            if descriptors[offset:offset + 1] != b"\r"
                        ]
                if len(header) == 32:
                    entry["dbf_rows"] = int.from_bytes(header[4:8], "little")
            if headers and info.filename.lower().endswith(".xlsx"):
                with _members(archive, {"xlsx": info.filename}) as paths:
                    book = load_workbook(paths["xlsx"], read_only=True, data_only=True)
                    try:
                        entry["sheets"] = {
                            sheet.title: [str(cell).strip() if cell is not None else ""
                                          for cell in next(sheet.iter_rows(values_only=True), ())]
                            for sheet in book.worksheets
                        }
                    finally:
                        book.close()
            members.append(entry)
    return members

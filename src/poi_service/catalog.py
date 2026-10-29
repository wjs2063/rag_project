"""Versioned category catalog and explicit visit-purpose allowlist."""

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


@dataclass(frozen=True)
class Category:
    id: str
    name: str
    parent_id: str | None
    allowed: bool


class Catalog:
    def __init__(self, categories: dict[str, Category], version: str):
        if not categories or not version:
            raise ValueError("catalog requires categories and a version")
        self.categories = categories
        self.version = version
        self.children: dict[str, set[str]] = {key: set() for key in categories}
        for category in categories.values():
            if category.parent_id:
                if category.parent_id not in categories:
                    raise ValueError(f"unknown category parent: {category.parent_id}")
                self.children[category.parent_id].add(category.id)

    @classmethod
    def from_csv(cls, path: Path, version: str) -> "Catalog":
        categories: dict[str, Category] = {}
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            if not {"id", "name", "parent_id", "allowed"} <= set(reader.fieldnames or []):
                raise ValueError("catalog CSV needs id,name,parent_id,allowed")
            for row in reader:
                key = row["id"].strip()
                if not key or key in categories:
                    raise ValueError(f"empty or duplicate category: {key}")
                flag = row["allowed"].strip().lower()
                if flag not in {"true", "false", "1", "0"}:
                    raise ValueError(f"invalid allowed value for {key}")
                categories[key] = Category(
                    key, row["name"].strip(), row["parent_id"].strip() or None, flag in {"true", "1"}
                )
        if not any(item.allowed for item in categories.values()):
            raise ValueError("catalog has no explicitly allowed category")
        return cls(categories, version)

    def descendants(self, key: str) -> set[str]:
        if key not in self.categories:
            raise KeyError(key)
        result = {key}
        pending = [key]
        while pending:
            for child in self.children[pending.pop()]:
                result.add(child)
                pending.append(child)
        return result

    def allowed_descendants(self, keys: list[str]) -> list[str]:
        selected: set[str] = set()
        for key in keys:
            selected.update(self.descendants(key))
        return sorted(key for key in selected if self.categories[key].allowed)

    def as_tree(self) -> dict:
        return {
            "version": self.version,
            "categories": [
                {"id": c.id, "name": c.name, "parent_id": c.parent_id, "allowed": c.allowed}
                for c in self.categories.values()
            ],
        }


def catalog_path(manifest_path: Path) -> tuple[Path, str]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    return manifest_path.parent / manifest["catalog_csv"], manifest["catalog_version"]


def build_catalog_from_xlsx(archive: Path, member: str, policy_file: Path, output: Path) -> dict:
    """Extract the source workbook's 16th-column codes under an explicit visit policy."""
    policy = json.loads(policy_file.read_text(encoding="utf-8"))
    with ZipFile(archive) as source:
        workbook_data = source.read(member)
    book = load_workbook(io.BytesIO(workbook_data), read_only=True, data_only=True)
    try:
        sheet = next(sheet for sheet in book.worksheets if "2,331" in sheet.title)
        category_rows = []
        seen = set()
        for row in sheet.iter_rows(min_row=3, values_only=True):
            code = str(row[15]).zfill(12) if row[15] is not None else ""
            if not re.fullmatch(r"[0-9]{12}", code) or code in seen:
                raise ValueError(f"invalid or duplicate source category code: {code}")
            seen.add(code)
            top, subgroup = str(row[0]), str(row[2])
            name = next(
                (str(row[index]).strip() for index in (10, 8, 6, 4, 2, 0)
                 if row[index] not in (None, "-", "")), ""
            )
            allowed = (
                top in policy["allow_top"]
                or subgroup in policy["allow_sub"].get(top, [])
                or any(code.startswith(prefix) for prefix in policy["allow_code_prefixes"])
            ) and code not in policy.get("deny_codes", [])
            category_rows.append((code, name, f"top_{code[0]}", allowed))
    finally:
        book.close()
    if len(category_rows) != policy["expected_source_codes"]:
        raise ValueError(f"expected {policy['expected_source_codes']} source codes, found {len(category_rows)}")
    top_names = {"1": "기관명", "2": "상호명", "4": "시설", "5": "지명"}
    if {code[0] for code, *_ in category_rows} != set(top_names):
        raise ValueError("unexpected top-level source category code")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(["id", "name", "parent_id", "allowed"])
        writer.writerows((f"top_{key}", name, "", "false") for key, name in top_names.items())
        writer.writerows((code, name, parent, str(allowed).lower())
                         for code, name, parent, allowed in category_rows)
    return {
        "version": policy["version"],
        "source_codes": len(category_rows),
        "allowed_codes": sum(allowed for *_, allowed in category_rows),
        "output": str(output),
    }
